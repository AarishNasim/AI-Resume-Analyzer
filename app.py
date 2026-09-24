import io, os, sqlite3
from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, send_file
from flask_cors import CORS
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from services.resume_parser import extract_text
from services.skill_extractor import extract_skills
from services.matcher import match_jobs
from services.template_engine import TEMPLATES, get_template, template_choices

BASE = os.path.dirname(os.path.abspath(__file__))
APP_ENV = os.environ.get('APP_ENV', 'development').lower()
DB = os.environ.get('DATABASE_PATH', os.path.join(BASE, 'resume_analyzer.db'))
UPLOADS = os.environ.get('UPLOADS_DIR', os.path.join(BASE, 'uploads'))
os.makedirs(UPLOADS, exist_ok=True)

app = Flask(__name__)
secret_key = os.environ.get('SECRET_KEY')
if APP_ENV == 'production' and not secret_key:
    raise RuntimeError('SECRET_KEY must be set when APP_ENV=production')
app.secret_key = secret_key or 'dev-change-this-secret-key'
app.config['MAX_CONTENT_LENGTH'] = 5 * 1024 * 1024
ALLOWED = {'pdf', 'docx'}

configured_origins = [origin.strip() for origin in os.environ.get('CORS_ORIGINS', '').split(',') if origin.strip()]
frontend_url = os.environ.get('FRONTEND_URL')
if frontend_url:
    configured_origins.append(frontend_url.rstrip('/'))
if not configured_origins and APP_ENV != 'production':
    configured_origins = ['http://127.0.0.1:5000', 'http://localhost:5000']
CORS(app, resources={r'/*': {'origins': configured_origins}}, supports_credentials=True)

RESUME_TEMPLATES = {
    'classic': ('Classic Professional', 'Traditional, polished, and easy for ATS systems to scan.', 'CLASSIC PROFESSIONAL RESUME'),
    'modern': ('Modern Minimal', 'Clean hierarchy with a confident, contemporary feel.', 'MODERN MINIMAL RESUME'),
    'executive': ('Executive Impact', 'Space for leadership wins, scope, and measurable outcomes.', 'EXECUTIVE IMPACT RESUME'),
    'technical': ('Technical Builder', 'Skills-forward structure for engineering and technical roles.', 'TECHNICAL BUILDER RESUME'),
    'creative': ('Creative Portfolio', 'A flexible format for designers, writers, and brand makers.', 'CREATIVE PORTFOLIO RESUME'),
    'student': ('Early Career', 'Guided sections for students, interns, and career starters.', 'EARLY CAREER RESUME'),
}

def db():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    return con

def init_db():
    con = db()
    con.executescript('''
    CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, email TEXT UNIQUE NOT NULL, password TEXT NOT NULL, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS resumes(id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, filename TEXT, resume_text TEXT, skills TEXT, score REAL, uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY(user_id) REFERENCES users(id));
    CREATE TABLE IF NOT EXISTS jobs(id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, company TEXT NOT NULL, description TEXT NOT NULL, required_skills TEXT NOT NULL, location TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS applications(id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, job_id INTEGER, match_percentage REAL, status TEXT DEFAULT 'Saved', applied_at TEXT DEFAULT CURRENT_TIMESTAMP);
    ''')
    columns = {row[1] for row in con.execute('PRAGMA table_info(resumes)').fetchall()}
    if 'source_path' not in columns: con.execute('ALTER TABLE resumes ADD COLUMN source_path TEXT')
    if 'source_type' not in columns: con.execute('ALTER TABLE resumes ADD COLUMN source_type TEXT')
    if 'template_key' not in columns: con.execute('ALTER TABLE resumes ADD COLUMN template_key TEXT')
    if con.execute('SELECT COUNT(*) FROM jobs').fetchone()[0] == 0:
        jobs = [
        ('Python Developer','TechNova','Build Flask APIs and backend services.','Python, Flask, SQL, REST API, Git','Noida / Remote'),
        ('Data Analyst','DataBridge','Analyze business data and create reports.','Python, SQL, Pandas, Excel, Statistics','Delhi'),
        ('ML Intern','AI Labs','Work on machine learning and NLP experiments.','Python, Machine Learning, Pandas, Scikit-learn, NLP','Jaipur / Remote'),
        ('Frontend Developer','WebCraft','Create responsive web interfaces.','HTML, CSS, JavaScript, React, Git','Hyderabad'),
        ('Full Stack Developer','CodeWorks','Develop frontend and backend features.','Python, JavaScript, Flask, HTML, CSS, SQL','Noida')]
        con.executemany('INSERT INTO jobs(title,company,description,required_skills,location) VALUES(?,?,?,?,?)', jobs)
    con.commit(); con.close()

def ats_score(text):
    content = text.lower()
    skills = extract_skills(text)
    sections = ['experience', 'education', 'skills']
    section_score = sum(section in content for section in sections) / len(sections) * 30
    skill_score = min(len(skills) / 10, 1) * 35
    length_score = 20 if 250 <= len(text.split()) <= 900 else 10 if len(text.split()) >= 120 else 0
    contact_score = 15 if any(marker in content for marker in ['@', 'phone', 'linkedin']) else 0
    return round(min(section_score + skill_score + length_score + contact_score, 100))

def ai_improve_resume(text, target_role):
    target_role = target_role.strip() or 'your target role'
    replacements = {'responsible for': 'led', 'worked on': 'delivered', 'helped with': 'supported', 'made': 'created'}
    lines = text.splitlines()
    improved = []
    role_added = False
    for line in lines:
        updated = line
        if line.strip() and not line.strip().isupper() and not role_added:
            updated = f'{line.rstrip()} | Target role: {target_role}'
            role_added = True
        for old, new in replacements.items():
            updated = updated.replace(old, new).replace(old.title(), new.title())
        improved.append(updated)
    return '\n'.join(improved)

def update_docx_in_place(source_path, edited_text):
    from docx import Document
    document = Document(source_path)
    original_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
    edited_lines = [line for line in edited_text.splitlines() if line.strip()]
    for paragraph, value in zip(original_paragraphs, edited_lines):
        if paragraph.runs:
            paragraph.runs[0].text = value
            for run in paragraph.runs[1:]: run.text = ''
        else:
            paragraph.add_run(value)
    document.save(source_path)

def template_text(template_key):
    title = RESUME_TEMPLATES[template_key][2]
    return f'''{title}\n\nFULL NAME\nEmail | Phone | LinkedIn | City, Country\n\nPROFESSIONAL SUMMARY\nWrite a focused 2-3 line summary for your target role.\n\nEXPERIENCE\nJob Title | Company | Dates\n- Describe an achievement with a measurable result.\n- Add the tools, scope, and impact of your work.\n\nEDUCATION\nDegree | Institution | Graduation Year\n\nSKILLS\nAdd relevant technical and professional skills.\n'''

def builder_text(form):
    sections = [
        ('PROFESSIONAL SUMMARY', form.get('summary', '')),
        ('EXPERIENCE', form.get('experience', '')),
        ('EDUCATION', form.get('education', '')),
        ('PROJECTS', form.get('projects', '')),
        ('SKILLS', form.get('skills', '')),
    ]
    personal = ' | '.join(value for value in [form.get('full_name', ''), form.get('email', ''), form.get('phone', ''), form.get('links', '')] if value)
    return '\n\n'.join([personal] + [f'{title}\n{value}' for title, value in sections if value.strip()])

def builder_values(resume):
    text = resume['resume_text'] if resume else ''
    return {'full_name': '', 'email': '', 'phone': '', 'links': '', 'summary': text, 'experience': '', 'education': '', 'projects': '', 'skills': resume['skills'] if resume else ''}

@app.route('/')
def home(): return render_template('index.html')

@app.route('/register', methods=['GET','POST'])
def register():
    if request.method == 'POST':
        name=request.form.get('name','').strip(); email=request.form.get('email','').strip().lower(); password=request.form.get('password','')
        if not name or not email or len(password)<6: flash('Name, email and password (6+ characters) are required.','error'); return render_template('register.html')
        try:
            con=db(); con.execute('INSERT INTO users(name,email,password) VALUES(?,?,?)',(name,email,generate_password_hash(password))); con.commit(); con.close()
            flash('Registration successful. Please login.','success'); return redirect(url_for('login'))
        except sqlite3.IntegrityError: flash('Email already registered.','error')
    return render_template('register.html')

@app.route('/login', methods=['GET','POST'])
def login():
    if request.method=='POST':
        email=request.form.get('email','').strip().lower(); password=request.form.get('password','')
        con=db(); user=con.execute('SELECT * FROM users WHERE email=?',(email,)).fetchone(); con.close()
        if user and check_password_hash(user['password'],password):
            session['user_id']=user['id']; session['name']=user['name']; return redirect(url_for('dashboard'))
        flash('Invalid email or password.','error')
    return render_template('login.html')

@app.route('/logout')
def logout(): session.clear(); return redirect(url_for('home'))

@app.route('/dashboard')
def dashboard():
    if 'user_id' not in session: return redirect(url_for('login'))
    con=db(); resume=con.execute('SELECT * FROM resumes WHERE user_id=? ORDER BY id DESC LIMIT 1',(session['user_id'],)).fetchone(); jobs=con.execute('SELECT * FROM jobs').fetchall(); con.close()
    results=match_jobs(resume['resume_text'], jobs) if resume else []
    ats = ats_score(resume['resume_text']) if resume else 0
    active_template = (get_template(resume['template_key']) if resume and resume['template_key'] else None) or TEMPLATES[0]
    feedback = {
        'keywords': 'Add role-specific keywords from the job description.' if ats < 75 else 'Your keyword coverage is in good shape.',
        'formatting': 'Use standard section headings and keep bullet points concise.' if ats < 85 else 'Your structure is easy for ATS systems to scan.',
        'impact': 'Start bullets with strong verbs and include measurable outcomes.' if ats < 90 else 'Your bullets show clear impact.',
    }
    return render_template('dashboard.html', name=session['name'], resume=resume, results=results, ats=ats, templates=RESUME_TEMPLATES, active_template=active_template, feedback=feedback)

@app.route('/upload', methods=['POST'])
def upload():
    if 'user_id' not in session: return redirect(url_for('login'))
    file=request.files.get('resume')
    if not file or not file.filename: flash('Please choose a PDF or DOCX resume.','error'); return redirect(url_for('dashboard'))
    ext=file.filename.rsplit('.',1)[-1].lower() if '.' in file.filename else ''
    if ext not in ALLOWED: flash('Only PDF and DOCX files are allowed.','error'); return redirect(url_for('dashboard'))
    filename=secure_filename(file.filename); path=os.path.join(UPLOADS, filename); file.save(path)
    try: text=extract_text(path)
    except Exception as e: flash(f'Could not read resume: {e}','error'); return redirect(url_for('dashboard'))
    if not text.strip(): flash('No readable text found in the resume.','error'); return redirect(url_for('dashboard'))
    skills=extract_skills(text)
    con=db(); con.execute('DELETE FROM resumes WHERE user_id=?',(session['user_id'],)); con.execute('INSERT INTO resumes(user_id,filename,resume_text,skills,source_path,source_type) VALUES(?,?,?,?,?,?)',(session['user_id'],filename,text,', '.join(skills),path,ext)); con.commit(); con.close()
    flash(f'Resume analyzed. Found {len(skills)} skills.','success'); return redirect(url_for('dashboard'))

@app.route('/update-resume', methods=['POST'])
def update_resume():
    if 'user_id' not in session: return redirect(url_for('login'))
    text = request.form.get('resume_text', '').strip()
    if not text:
        flash('Add some resume content before saving.', 'error')
        return redirect(url_for('dashboard'))
    con = db(); resume = con.execute('SELECT * FROM resumes WHERE user_id=? ORDER BY id DESC LIMIT 1', (session['user_id'],)).fetchone()
    if resume and resume['source_type'] == 'docx' and resume['source_path'] and os.path.exists(resume['source_path']):
        update_docx_in_place(resume['source_path'], text)
    con.execute('UPDATE resumes SET resume_text=?, skills=? WHERE user_id=?', (text, ', '.join(extract_skills(text)), session['user_id'])); con.commit(); con.close()
    flash('Resume updated and ready to download.', 'success')
    return redirect(url_for('dashboard'))

@app.route('/templates')
def template_gallery():
    if 'user_id' not in session: return redirect(url_for('login'))
    return render_template('template_gallery.html', template_groups=template_choices(), templates=TEMPLATES)

@app.route('/builder/<template_id>', methods=['GET', 'POST'])
def builder(template_id):
    if 'user_id' not in session: return redirect(url_for('login'))
    template = get_template(template_id)
    if not template: return redirect(url_for('template_gallery'))
    con = db(); resume = con.execute('SELECT * FROM resumes WHERE user_id=? ORDER BY id DESC LIMIT 1', (session['user_id'],)).fetchone()
    if request.method == 'POST':
        text = builder_text(request.form)
        if not text.strip():
            flash('Add at least your name or one resume section before saving.', 'error')
            return render_template('builder.html', template=template, values=request.form, ats=0)
        filename = f'{template_id}_resume.docx'
        skills = extract_skills(text)
        if resume:
            con.execute('UPDATE resumes SET filename=?, resume_text=?, skills=?, template_key=?, source_path=NULL, source_type=? WHERE user_id=?', (filename, text, ', '.join(skills), template_id, 'template', session['user_id']))
        else:
            con.execute('INSERT INTO resumes(user_id,filename,resume_text,skills,template_key,source_type) VALUES(?,?,?,?,?,?)', (session['user_id'], filename, text, ', '.join(skills), template_id, 'template'))
        con.commit(); con.close()
        flash('Resume saved with your selected template.', 'success')
        return redirect(url_for('dashboard'))
    con.close()
    values = builder_values(resume)
    return render_template('builder.html', template=template, values=values, ats=ats_score(resume['resume_text']) if resume else 0)

@app.route('/ai-improve', methods=['POST'])
def ai_improve():
    if 'user_id' not in session: return jsonify({'error': 'Please log in first.'}), 401
    payload = request.get_json(silent=True) or {}
    text = payload.get('resume_text', '').strip()
    if not text: return jsonify({'error': 'Add resume text first.'}), 400
    return jsonify({'resume_text': ai_improve_resume(text, payload.get('target_role', ''))})

@app.route('/download-resume')
def download_resume():
    if 'user_id' not in session: return redirect(url_for('login'))
    con = db(); resume = con.execute('SELECT * FROM resumes WHERE user_id=? ORDER BY id DESC LIMIT 1', (session['user_id'],)).fetchone(); con.close()
    if not resume:
        flash('Upload a resume before downloading.', 'error')
        return redirect(url_for('dashboard'))
    if resume['source_type'] == 'docx' and resume['source_path'] and os.path.exists(resume['source_path']):
        return send_file(resume['source_path'], as_attachment=True, download_name='updated_resume.docx', mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    from docx import Document
    document = Document()
    for line in resume['resume_text'].splitlines():
        document.add_paragraph(line)
    output = io.BytesIO(); document.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name='updated_resume.docx', mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@app.route('/download-pdf')
def download_pdf():
    if 'user_id' not in session: return redirect(url_for('login'))
    con = db(); resume = con.execute('SELECT * FROM resumes WHERE user_id=? ORDER BY id DESC LIMIT 1', (session['user_id'],)).fetchone(); con.close()
    if not resume: return redirect(url_for('template_gallery'))
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    output = io.BytesIO(); document = SimpleDocTemplate(output, pagesize=LETTER, rightMargin=.7*inch, leftMargin=.7*inch, topMargin=.65*inch, bottomMargin=.65*inch)
    styles = getSampleStyleSheet(); body = ParagraphStyle('ResumeBody', parent=styles['BodyText'], fontName='Helvetica', fontSize=9.5, leading=13, spaceAfter=6); heading = ParagraphStyle('ResumeHeading', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
    story = []
    for line in resume['resume_text'].splitlines():
        if not line.strip(): story.append(Spacer(1, 3)); continue
        safe_line = line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        style = heading if line.strip().isupper() else body
        story.append(Paragraph(safe_line, style))
    document.build(story); output.seek(0)
    return send_file(output, as_attachment=True, download_name='resume_ats_ready.pdf', mimetype='application/pdf')

@app.route('/new-resume/<template_key>')
def new_resume(template_key):
    if 'user_id' not in session: return redirect(url_for('login'))
    if template_key not in RESUME_TEMPLATES: return redirect(url_for('dashboard'))
    text = template_text(template_key)
    con = db(); con.execute('DELETE FROM resumes WHERE user_id=?', (session['user_id'],)); con.execute('INSERT INTO resumes(user_id,filename,resume_text,skills,template_key,source_type) VALUES(?,?,?,?,?,?)', (session['user_id'], f'{template_key}_resume.docx', text, ', '.join(extract_skills(text)), template_key, 'template')); con.commit(); con.close()
    flash(f'{RESUME_TEMPLATES[template_key][0]} template opened. Replace the placeholder data and save.', 'success')
    return redirect(url_for('dashboard'))

@app.route('/save-job/<int:job_id>', methods=['POST'])
def save_job(job_id):
    if 'user_id' not in session: return redirect(url_for('login'))
    con=db(); resume=con.execute('SELECT * FROM resumes WHERE user_id=? ORDER BY id DESC LIMIT 1',(session['user_id'],)).fetchone(); job=con.execute('SELECT * FROM jobs WHERE id=?',(job_id,)).fetchone()
    if resume and job:
        score=next((x['score'] for x in match_jobs(resume['resume_text'],[job]) if x['id']==job_id),0)
        con.execute('INSERT INTO applications(user_id,job_id,match_percentage) VALUES(?,?,?)',(session['user_id'],job_id,score)); con.commit()
    con.close(); flash('Job saved.','success'); return redirect(url_for('dashboard'))

if __name__=='__main__':
    port = int(os.environ.get('PORT', 5000))
    init_db()
    app.run(host='0.0.0.0', port=port)
else: init_db()
